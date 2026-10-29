"""Multi-format product catalogue export engine (XLSX, CSV, JSON)."""

import csv
import io
import json
import re
from typing import Any, Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.catalogue.models import ProductRecord
from app.catalogue.templates import ProductTemplate, get_template_by_id


def export_catalogue(
    products: List[ProductRecord],
    template_id: str = "standard_ecommerce",
    file_format: str = "xlsx",
    custom_filename: str | None = None,
    org_id: str = "default",
    template_override: ProductTemplate | None = None,
) -> Tuple[bytes, str, str]:
    """Export product records in the specified format formatted to the target template schema."""
    template = template_override or get_template_by_id(template_id, org_id=org_id)
    if not template:
        template = get_template_by_id("standard_ecommerce")

    fmt = file_format.lower().strip()
    if fmt not in ("xlsx", "csv", "json"):
        fmt = "xlsx"

    # Build mapped rows according to template field order
    mapped_rows: List[Dict[str, Any]] = []
    field_columns = [f.name for f in template.fields]

    for p in products:
        row: Dict[str, Any] = {}
        p_dict = p.model_dump()

        for f in template.fields:
            col_name = f.name
            mapped_key = f.mapped_to or f.name

            # Special computed fields
            if mapped_key == "_shopify_handle":
                handle = re.sub(r"[^a-zA-Z0-9]+", "-", (p.name or "product").lower()).strip("-")
                val = handle or f"product-{p.sku.lower()}"
            elif mapped_key in p_dict:
                val = p_dict[mapped_key]
            elif mapped_key in p.custom_attributes:
                val = p.custom_attributes[mapped_key]
            else:
                val = f.default_value

            # Format formatting by data type
            if f.data_type == "tags":
                if isinstance(val, list):
                    val = ", ".join(str(x) for x in val if str(x).strip())
                elif val is None:
                    val = ""
            elif f.data_type in ("number", "integer"):
                if val is None or val == "":
                    val = f.default_value if f.default_value is not None else 0
                else:
                    try:
                        val = int(val) if f.data_type == "integer" else round(float(val), 2)
                    except (ValueError, TypeError):
                        val = 0
            elif val is None:
                val = f.default_value if f.default_value is not None else ""

            row[col_name] = val

        mapped_rows.append(row)

    base_name = custom_filename or f"rivu_{template.platform}_catalog"
    base_name = re.sub(r"[^a-zA-Z0-9_-]", "_", base_name).strip("_")

    if fmt == "xlsx":
        data = _generate_xlsx(field_columns, mapped_rows, template.name)
        filename = f"{base_name}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "csv":
        data = _generate_csv(field_columns, mapped_rows)
        filename = f"{base_name}.csv"
        media_type = "text/csv; charset=utf-8"
    else:  # json
        data = _generate_json(mapped_rows)
        filename = f"{base_name}.json"
        media_type = "application/json"

    return data, filename, media_type


def generate_export_preview(
    products: List[ProductRecord],
    template_id: str = "standard_ecommerce",
    file_format: str = "xlsx",
    org_id: str = "default",
    template_override: ProductTemplate | None = None,
) -> Dict[str, Any]:
    """Generate a lightweight preview of the export without building the full binary file."""
    template = template_override or get_template_by_id(template_id, org_id=org_id)
    if not template:
        template = get_template_by_id("standard_ecommerce")

    field_columns = [f.name for f in template.fields]
    sample_products = products[:5]

    sample_rows: List[Dict[str, Any]] = []
    for p in sample_products:
        row: Dict[str, Any] = {}
        p_dict = p.model_dump()
        for f in template.fields:
            mapped_key = f.mapped_to or f.name
            if mapped_key == "_shopify_handle":
                handle = re.sub(r"[^a-zA-Z0-9]+", "-", (p.name or "product").lower()).strip("-")
                row[f.name] = handle or f"product-{p.sku.lower()}"
            elif mapped_key in p_dict:
                val = p_dict[mapped_key]
                if isinstance(val, list):
                    val = ", ".join(str(x) for x in val)
                row[f.name] = val
            elif mapped_key in p.custom_attributes:
                row[f.name] = p.custom_attributes[mapped_key]
            else:
                row[f.name] = f.default_value if f.default_value is not None else ""
        sample_rows.append(row)

    valid_count = sum(1 for p in products if p.validation_status == "valid")
    error_count = sum(1 for p in products if p.validation_status == "error")

    return {
        "template_name": template.name,
        "format": file_format.lower(),
        "total_records": len(products),
        "valid_records": valid_count,
        "error_records": error_count,
        "output_columns": field_columns,
        "sample_rows": sample_rows,
    }


def _generate_xlsx(headers: List[str], rows: List[Dict[str, Any]], sheet_title: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r"[:\\/?*\[\]]", "", sheet_title)[:30] or "Products"

    # Styling: Rivu Dark Purple Theme
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F1435", end_color="1F1435", fill_type="solid")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=False)

    data_font = Font(name="Segoe UI", size=10, color="222222")
    data_align = Alignment(vertical="center")

    # Write Header
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 28

    # Write Rows
    for r_idx, row in enumerate(rows, start=2):
        row_vals = [row.get(h, "") for h in headers]
        ws.append(row_vals)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.alignment = data_align
        ws.row_dimensions[r_idx].height = 20

    # Auto-fit column widths
    for col_idx, h in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(h))
        for r in range(1, min(len(rows) + 2, 50)):
            val = ws.cell(row=r, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)[:50]))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_csv(headers: List[str], rows: List[Dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=headers,
        quoting=csv.QUOTE_MINIMAL,
        lineterminator="\r\n",
    )
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    # Return as UTF-8 with BOM for native Excel compatibility
    return buf.getvalue().encode("utf-8-sig")


def _generate_json(rows: List[Dict[str, Any]]) -> bytes:
    content = json.dumps(rows, indent=2, ensure_ascii=False)
    return content.encode("utf-8")
