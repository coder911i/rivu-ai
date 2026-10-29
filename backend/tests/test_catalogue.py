"""Unit and integration tests for the Catalogue-to-Product Data Entry solution."""

import io
import json
import openpyxl
import pytest
from httpx import AsyncClient, ASGITransport
from pypdf import PdfWriter

from app.main import app
from app.catalogue.parser import parse_catalogue_file, detect_catalogue_format
from app.catalogue.mapping import infer_field_mappings, apply_field_mappings, _apply_transformation
from app.catalogue.validator import validate_product_records
from app.catalogue.templates import (
    get_all_templates, get_template_by_id, save_custom_template, delete_custom_template,
    ProductTemplate, TemplateField
)
from app.catalogue.exporter import export_catalogue, generate_export_preview
from app.catalogue.models import ProductRecord, FieldMapping
from app.core.security import create_access_token


# ── Test 1: Parser Multi-format Support ──────────────────────────

def test_detect_catalogue_format():
    assert detect_catalogue_format("products.pdf") == "pdf"
    assert detect_catalogue_format("catalog.xlsx") == "xlsx"
    assert detect_catalogue_format("items.xls") == "xls"
    assert detect_catalogue_format("feed.csv") == "csv"
    assert detect_catalogue_format("inventory.json") == "json"

    with pytest.raises(Exception):
        detect_catalogue_format("malicious.exe")


def test_parse_csv_catalogue():
    csv_bytes = (
        "Item SKU,Product Description,Retail MSRP,Available Qty\n"
        "SKU-1001,Wireless Ergonomic Mouse,$29.99,150\n"
        "SKU-1002,Mechanical Gaming Keyboard,$89.50,75\n"
    ).encode("utf-8")

    parsed = parse_catalogue_file(csv_bytes, "supplier_feed.csv")
    assert parsed["file_format"] == "csv"
    assert parsed["row_count"] == 2
    assert "Item SKU" in parsed["columns"]
    assert "Retail MSRP" in parsed["columns"]
    assert parsed["raw_records"][0]["Item SKU"] == "SKU-1001"
    assert parsed["raw_records"][0]["Retail MSRP"] == "$29.99"


def test_parse_excel_catalogue_multi_sheet():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Consumer Electronics"
    ws1.append(["Part Number", "Title", "Unit Cost", "Stock"])
    ws1.append(["CE-01", "USB-C Fast Charger 65W", "18.50", "300"])

    ws2 = wb.create_sheet("Audio")
    ws2.append(["Part Number", "Title", "Unit Cost", "Stock"])
    ws2.append(["AU-01", "Noise Cancelling Headphones", "120.00", "50"])

    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()

    # Parse default sheet
    parsed1 = parse_catalogue_file(excel_bytes, "catalogue.xlsx")
    assert parsed1["file_format"] == "xlsx"
    assert "Consumer Electronics" in parsed1["sheet_names"]
    assert "Audio" in parsed1["sheet_names"]
    assert parsed1["row_count"] == 1
    assert parsed1["raw_records"][0]["Part Number"] == "CE-01"

    # Parse specific sheet
    parsed2 = parse_catalogue_file(excel_bytes, "catalogue.xlsx", sheet_name="Audio")
    assert parsed2["row_count"] == 1
    assert parsed2["raw_records"][0]["Part Number"] == "AU-01"


def test_parse_json_catalogue():
    json_bytes = json.dumps({
        "products": [
            {
                "sku": "J-01",
                "name": "Smart Water Bottle",
                "price": 34.99,
                "specs": {"capacity": "750ml", "color": "Midnight Black"}
            },
            {
                "sku": "J-02",
                "name": "Stainless Steel Mug",
                "price": 19.99,
                "specs": {"capacity": "350ml", "color": "Silver"}
            }
        ]
    }).encode("utf-8")

    parsed = parse_catalogue_file(json_bytes, "catalog.json")
    assert parsed["file_format"] == "json"
    assert parsed["row_count"] == 2
    assert "sku" in parsed["columns"]
    assert "specs_color" in parsed["columns"]
    assert parsed["raw_records"][0]["specs_color"] == "Midnight Black"


def test_parse_pdf_catalogue():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(50, 750, "PartNum    ItemName                       Price     Qty")
    c.drawString(50, 720, "SKU-9901   Industrial Ball Bearing        $14.25    200")
    c.drawString(50, 690, "SKU-9902   High-Pressure Hydraulic Hose   $48.50    50")
    c.save()
    pdf_bytes = buf.getvalue()

    parsed = parse_catalogue_file(pdf_bytes, "supplier_list.pdf")
    assert parsed["file_format"] == "pdf"
    assert parsed["total_pages"] == 1
    assert parsed["row_count"] >= 1
    assert any("SKU-9901" in str(v) for row in parsed["raw_records"] for v in row.values())


# ── Test 2: Field Mapping & Auto-Inference ────────────────────────

def test_infer_field_mappings():
    columns = ["Item Part No", "Product Description", "MSRP Price", "Quantity On Hand", "Vendor Name"]
    sample_rows = [
        {
            "Item Part No": "PRT-8821",
            "Product Description": "Titanium Watch Strap 22mm",
            "MSRP Price": "$125.00",
            "Quantity On Hand": "45",
            "Vendor Name": "Horology Labs",
        }
    ]

    mappings = infer_field_mappings(columns, sample_rows)
    mapped_targets = {m.source_column: m.target_field for m in mappings}

    assert mapped_targets.get("Item Part No") == "sku"
    assert mapped_targets.get("Product Description") == "name"
    assert mapped_targets.get("MSRP Price") == "price"
    assert mapped_targets.get("Quantity On Hand") == "stock_quantity"
    assert mapped_targets.get("Vendor Name") == "brand"


def test_field_transformations():
    assert _apply_transformation("  hello world  ", "trim") == "hello world"
    assert _apply_transformation("sku_123_abc", "uppercase") == "SKU_123_ABC"
    assert _apply_transformation("  $1,299.95 USD ", "currency_to_decimal") == 1299.95
    assert _apply_transformation("42.8", "to_integer") == 42
    assert _apply_transformation("Smart Watch Pro Max!", "slugify") == "smart-watch-pro-max"


def test_apply_field_mappings():
    raw_records = [
        {"Part #": "A-101", "Item": "Running Shoes", "Retail": "$79.99", "Stock": "25", "Material": "Mesh"},
        {"Part #": "B-202", "Item": "Trail Boots", "Retail": "€119.50", "Stock": "10", "Material": "Leather"},
    ]
    mappings = [
        FieldMapping(source_column="Part #", target_field="sku"),
        FieldMapping(source_column="Item", target_field="name"),
        FieldMapping(source_column="Retail", target_field="price", transform="currency_to_decimal"),
        FieldMapping(source_column="Stock", target_field="stock_quantity", transform="to_integer"),
        FieldMapping(source_column="Material", target_field="custom:Material"),
    ]

    products = apply_field_mappings(raw_records, mappings)
    assert len(products) == 2
    assert products[0].sku == "A-101"
    assert products[0].name == "Running Shoes"
    assert products[0].price == 79.99
    assert products[0].stock_quantity == 25
    assert products[0].custom_attributes.get("Material") == "Mesh"


def test_missing_identifiers_are_not_fabricated():
    products = apply_field_mappings([{"Supplier note": "A real source fact"}], [])
    assert products[0].name == ""
    assert products[0].sku == ""
    validated, summary = validate_product_records(products)
    assert summary.error_records == 1
    assert {issue.field for issue in validated[0].validation_issues} >= {"name", "sku"}


def test_ai_fallback_does_not_invent_product_facts():
    from app.catalogue.ai_enrich import _generate_rule_based_fallback

    product = ProductRecord(id="p1", name="Steel Thermos", sku="TH-1", price=10)
    suggestion = _generate_rule_based_fallback(product, ["description", "category", "seo", "tags"])
    assert suggestion.suggested_description is None
    assert suggestion.suggested_category is None
    assert suggestion.suggested_seo_description == "Steel Thermos"


def test_custom_template_required_fields_are_validated():
    template = ProductTemplate(
        id="custom_required", name="Required", description="Required field test",
        fields=[TemplateField(name="WarrantyCode", label="Warranty", required=True, mapped_to="custom:WarrantyCode")],
    )
    products, summary = validate_product_records(
        [ProductRecord(id="p1", name="Widget", sku="W-1", price=10)], template,
    )
    assert summary.error_records == 1
    assert any(issue.field == "custom:WarrantyCode" for issue in products[0].validation_issues)


# ── Test 3: Validation Engine ────────────────────────────────────

def test_validator_detects_errors_and_duplicates():
    products = [
        ProductRecord(id="1", name="Valid Widget", sku="W-01", price=19.99, stock_quantity=10),
        ProductRecord(id="2", name="", sku="W-02", price=25.00),  # Missing name
        ProductRecord(id="3", name="Invalid Price", sku="W-03", price=-5.00),  # Negative price
        ProductRecord(id="4", name="Duplicate A", sku="W-DUP", price=10.00),  # Duplicate SKU
        ProductRecord(id="5", name="Duplicate B", sku="W-DUP", price=12.00),  # Duplicate SKU
        ProductRecord(id="6", name="Zero Price Item", sku="W-04", price=0.00),  # Warning: price is 0
    ]

    validated, summary = validate_product_records(products)
    assert summary.total_records == 6
    assert summary.error_records >= 4
    assert summary.warning_records >= 1
    assert summary.valid_records == 1

    # Check product 1 is valid
    assert validated[0].validation_status == "valid"
    # Check product 2 has name error
    assert any(i.field == "name" for i in validated[1].validation_issues)
    # Check product 3 has price error
    assert any(i.field == "price" and i.severity == "error" for i in validated[2].validation_issues)
    # Check duplicate SKUs flagged
    assert any(i.issue_type == "duplicate_sku" for i in validated[3].validation_issues)
    assert any(i.issue_type == "duplicate_sku" for i in validated[4].validation_issues)


# ── Test 4: Dynamic Templates ────────────────────────────────────

def test_built_in_and_custom_templates():
    templates = get_all_templates(org_id="test_org")
    ids = [t.id for t in templates]
    assert "standard_ecommerce" in ids
    assert "shopify" in ids
    assert "woocommerce" in ids
    assert "erp_inventory" in ids

    # Create custom template
    custom = ProductTemplate(
        id="custom_b2b",
        name="B2B Wholesale Portal",
        description="Custom wholesale schema",
        platform="custom",
        fields=[
            TemplateField(name="WholesaleSKU", label="SKU", data_type="string", required=True, mapped_to="sku"),
            TemplateField(name="ProductName", label="Name", data_type="string", required=True, mapped_to="name"),
            TemplateField(name="Tier1Price", label="Tier 1 Price", data_type="number", required=True, mapped_to="price"),
        ]
    )
    saved = save_custom_template("test_org", custom)
    assert saved.id == "custom_b2b"

    fetched = get_template_by_id("custom_b2b", org_id="test_org")
    assert fetched is not None
    assert fetched.name == "B2B Wholesale Portal"

    # Delete custom template
    deleted = delete_custom_template("test_org", "custom_b2b")
    assert deleted is True
    assert get_template_by_id("custom_b2b", org_id="test_org") is None


# ── Test 5: Multi-Format Exporter ────────────────────────────────

def test_export_xlsx_csv_json():
    products = [
        ProductRecord(
            id="p1",
            name="Ergonomic Aluminum Keyboard Stand",
            sku="ACC-101",
            brand="ErgoDesk",
            category="Office Accessories",
            price=49.95,
            sale_price=39.95,
            stock_quantity=80,
            currency="USD",
            weight=1.2,
            tags=["desk", "aluminum", "ergonomic"],
        ),
        ProductRecord(
            id="p2",
            name="Vertical Ergonomic Mouse Wireless",
            sku="ACC-102",
            brand="ErgoDesk",
            category="Peripherals",
            price=29.99,
            stock_quantity=120,
            currency="USD",
            weight=0.3,
            tags=["mouse", "wireless"],
        ),
    ]

    # 1. Export XLSX
    xlsx_bytes, xlsx_name, xlsx_mime = export_catalogue(products, template_id="shopify", file_format="xlsx")
    assert xlsx_name.endswith(".xlsx")
    assert len(xlsx_bytes) > 0
    # Verify openpyxl can read it back cleanly
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet = wb.active
    assert sheet.cell(row=1, column=1).value == "Handle"
    assert sheet.cell(row=2, column=2).value == "Ergonomic Aluminum Keyboard Stand"

    # 2. Export CSV
    csv_bytes, csv_name, csv_mime = export_catalogue(products, template_id="standard_ecommerce", file_format="csv")
    assert csv_name.endswith(".csv")
    decoded_csv = csv_bytes.decode("utf-8-sig")
    assert "sku,name,brand,category" in decoded_csv.lower()
    assert "ACC-101" in decoded_csv
    assert "ACC-102" in decoded_csv

    # 3. Export JSON
    json_bytes, json_name, json_mime = export_catalogue(products, template_id="erp_inventory", file_format="json")
    assert json_name.endswith(".json")
    parsed_json = json.loads(json_bytes.decode("utf-8"))
    assert len(parsed_json) == 2
    assert parsed_json[0]["ItemCode"] == "ACC-101"
    assert parsed_json[0]["ItemName"] == "Ergonomic Aluminum Keyboard Stand"
    assert parsed_json[0]["UnitPrice"] == 49.95

    # 4. Export Preview
    preview = generate_export_preview(products, template_id="shopify", file_format="xlsx")
    assert preview["total_records"] == 2
    assert "Handle" in preview["output_columns"]
    assert len(preview["sample_rows"]) == 2
